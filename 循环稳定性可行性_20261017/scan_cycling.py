"""Read-only screening of the current library; writes only this new assessment folder."""
from pathlib import Path
import json, re, hashlib, collections, sys
from pypdf import PdfReader
from docx import Document

sys.stdout.reconfigure(encoding='utf-8')
OUT = Path(__file__).resolve().parent
ROOT = OUT.parent
def read(p): return json.loads(p.read_text(encoding='utf-8-sig'))
def write(p, value): p.write_text(json.dumps(value, ensure_ascii=False, indent=2), encoding='utf-8')
audit = {r['id']: r for r in read(ROOT/'训练适用性审核_20261016/逐份训练用途审核.json')}
locations = read(ROOT/'文献整理_20261017/当前唯一文献位置.json')
old = {r['sha256']: r for r in read(ROOT/'训练适用性审核_20261016/正文证据定位.json')}
cycle = re.compile(r'cycl(?:ing|es|e\s+stability|ic\s+stability)|循环|siklus', re.I)
quant = re.compile(r'\d[\d,.\s]*\s*%|\d[\d,\s]*\s*cycles', re.I)
results=[]
for loc in locations:
    if not loc['in_active_library']: continue
    a=audit[loc['id']]
    pdf=Path(loc['local_pdf'])
    actual=hashlib.sha256(pdf.read_bytes()).hexdigest()
    assert actual==loc['sha256'], (loc['id'], 'source hash changed')
    cached=old.get(actual, {})
    cache=Path(cached.get('fulltext_file', 'missing'))
    if cache.is_file():
        chunks=re.split(r'\n===== PDF PAGE (\d+) =====\n',cache.read_text(encoding='utf-8-sig'))
        pages=[{'page':int(chunks[i]),'text':chunks[i+1]} for i in range(1,len(chunks),2)]
        assert len(pages)==a['pages']
        method='existing_text_locator_with_current_pdf_sha256_verified'
    else:
        pages=[{'page':i+1,'text':p.extract_text() or ''} for i,p in enumerate(PdfReader(pdf).pages)]
        method='pypdf_current_original'
    hits=[]
    for p in pages:
        txt=p['text']
        # Bibliographic mentions are not experimental observations.
        ref=re.search(r'(?im)^\s*(?:\d+\.?\s*)?References\s*$',txt)
        body=txt[:ref.start()] if ref else txt
        flat=re.sub(r'\s+',' ',body)
        spans=[]
        for m in cycle.finditer(flat):
            lo=max(0,m.start()-330); hi=min(len(flat),m.end()+550)
            if spans and lo<=spans[-1][1]: spans[-1][1]=hi
            else: spans.append([lo,hi])
        for lo,hi in spans:
            excerpt=flat[lo:hi]
            hits.append({'page':p['page'],'text':excerpt,'quantitative_nearby':bool(quant.search(excerpt))})
    item={'id':loc['id'],'title':a['title'],'doi':a.get('doi'),'year':a.get('year'),
          'category':loc['category'],'path':str(pdf),'sha256':actual,'text_method':method,
          'prior_restrictions':a['restrictions'],'hits':hits,'supplements':loc['supplements']}
    results.append(item)
    if loc['category'] in ['A','B','C']:
        (OUT/'locators').mkdir(exist_ok=True)
        (OUT/'locators'/f"{loc['id']}.txt").write_text(
            '\n'.join([f"{loc['id']} | {a['title']} | {a.get('doi')}",a['restrictions']]+[
                f"\nPDF p{h['page']} | numeric={h['quantitative_nearby']}\n{h['text']}" for h in hits]),encoding='utf-8')
        (OUT/'page_text').mkdir(exist_ok=True)
        write(OUT/'page_text'/f"{loc['id']}.json",pages)
write(OUT/'全库循环证据定位.json',results)
for loc in locations:
    if not loc['in_active_library']: continue
    for i,s in enumerate(loc['supplements']):
        p=Path(s)
        if p.suffix.lower()=='.pdf':
            pages=[{'page':j+1,'text':q.extract_text() or ''} for j,q in enumerate(PdfReader(p).pages)]
        else:
            d=Document(p)
            pages=[{'section':'docx_paragraphs_tables','text':'\n'.join([q.text for q in d.paragraphs]+[
                ' | '.join(c.text for c in row.cells) for t in d.tables for row in t.rows])}]
        write(OUT/'page_text'/f"{loc['id']}_SI{i+1}.json",pages)
summary=[{'id':x['id'],'category':x['category'],'pages':sorted({h['page'] for h in x['hits'] if h['quantitative_nearby']}),
          'hit_count':len(x['hits'])} for x in results]
write(OUT/'筛查索引.json',summary)
print(json.dumps({'scanned':len(results),'hashes_verified':len(results),'by_category':dict(collections.Counter(x['category'] for x in results)),
                 'AB_quantitative_locators':[s for s in summary if s['category'] in ['A','B']]},ensure_ascii=False,indent=2))
